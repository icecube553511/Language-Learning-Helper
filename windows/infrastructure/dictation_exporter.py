from __future__ import annotations

from pathlib import Path
from io import BytesIO
import zipfile

from application.ports import ExportedDictation
from domain.dictation import DictationDocument, DictationFormat


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


class DictationDocumentExporter:
    def export(self, document: DictationDocument, output_format: DictationFormat | None = None) -> ExportedDictation:
        if not document.rows:
            raise ValueError("不能导出空默写文档")
        output_format = output_format or next(iter(document.options.formats))
        if output_format is DictationFormat.DOCX:
            return ExportedDictation(self._docx(document), "dictation.docx", DOCX_MIME)
        return ExportedDictation(self._pdf(document), "dictation.pdf", PDF_MIME)

    def _docx(self, document: DictationDocument) -> bytes:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.shared import Inches, Mm, Pt

        doc = Document()
        section = doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = section.bottom_margin = Inches(0.55)
        section.left_margin = section.right_margin = Inches(0.65)
        for index, row in enumerate(document.rows, 1):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(f"{index}. {row.word}  {row.part_of_speech}  {row.meaning}")
            self._set_docx_font(run, document.options.font_size_pt)
            for offset, example in enumerate(row.examples):
                chinese = row.example_chinese[offset] if offset < len(row.example_chinese) else ""
                suffix = f"（{chinese}）" if chinese else ""
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(f"    e.g. {example}{suffix}")
                self._set_docx_font(run, document.options.font_size_pt)
        output = BytesIO()
        doc.save(output)
        return output.getvalue()

    @staticmethod
    def _set_docx_font(run, size: int) -> None:
        from docx.oxml.ns import qn
        from docx.shared import Pt

        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    def _pdf(self, document: DictationDocument) -> bytes:
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        font_name = self._register_pdf_font()
        style = ParagraphStyle(
            "Dictation", fontName=font_name, fontSize=document.options.font_size_pt,
            leading=document.options.font_size_pt * 1.55, alignment=TA_LEFT,
        )
        output = BytesIO()
        doc = SimpleDocTemplate(
            output, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
            topMargin=14 * mm, bottomMargin=14 * mm,
        )
        story = []
        for index, row in enumerate(document.rows, 1):
            text = f"{index}. {row.word}  {row.part_of_speech}  {row.meaning}"
            story.append(Paragraph(self._escape_pdf(text), style))
            for offset, example in enumerate(row.examples):
                chinese = row.example_chinese[offset] if offset < len(row.example_chinese) else ""
                suffix = f"（{chinese}）" if chinese else ""
                story.append(Paragraph(self._escape_pdf(f"    e.g. {example}{suffix}"), style))
            story.append(Spacer(1, document.options.font_size_pt * 0.5))
        doc.build(story)
        return output.getvalue()

    @staticmethod
    def _escape_pdf(value: str) -> str:
        return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    @staticmethod
    def _register_pdf_font() -> str:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        candidates = [
            Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/msyh.ttf"),
            Path("C:/Windows/Fonts/simhei.ttf"), Path("C:/Windows/Fonts/simsun.ttc"),
        ]
        for path in candidates:
            if path.is_file():
                try:
                    pdfmetrics.registerFont(TTFont("LexiTrackCJK", str(path)))
                    return "LexiTrackCJK"
                except Exception:
                    continue
        return "Helvetica"
